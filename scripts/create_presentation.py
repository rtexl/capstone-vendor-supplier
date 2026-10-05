from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "VendorLens_Presentation_deck.pptx"

# Brand palette: sober enterprise blues with teal/orange accents.
NAVY = RGBColor(13, 34, 58)
NAVY_2 = RGBColor(22, 51, 82)
BLUE = RGBColor(31, 99, 191)
TEAL = RGBColor(13, 148, 136)
VIOLET = RGBColor(109, 76, 181)
ORANGE = RGBColor(232, 109, 37)
GREEN = RGBColor(31, 137, 84)
RED = RGBColor(190, 58, 67)
INK = RGBColor(27, 42, 58)
MUTED = RGBColor(91, 108, 126)
LINE = RGBColor(211, 222, 233)
PALE = RGBColor(244, 248, 252)
PALE_BLUE = RGBColor(235, 244, 253)
PALE_TEAL = RGBColor(234, 249, 246)
PALE_ORANGE = RGBColor(255, 246, 237)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, *, line=None, rounded=False, kind=None):
    shape_kind = kind or (
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
        if rounded
        else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    )
    item = slide.shapes.add_shape(
        shape_kind, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    item.fill.solid()
    item.fill.fore_color.rgb = fill
    item.line.color.rgb = line or fill
    if rounded and item.adjustments:
        item.adjustments[0] = 0.1
    return item


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    *,
    size=18,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    font="Aptos",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, *, size=13, color=INK, gap=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(gap)
    return box


def footer(slide, page, label):
    rect(slide, 0, 7.08, 13.333, 0.42, NAVY)
    text(slide, "VendorLens AI", 0.66, 7.17, 2.2, 0.16, size=9, color=RGBColor(205, 220, 235), bold=True)
    text(slide, label, 4.1, 7.17, 5.1, 0.16, size=8.5, color=RGBColor(175, 198, 221), align=PP_ALIGN.CENTER)
    text(slide, f"{page:02d}", 12.0, 7.16, 0.65, 0.16, size=9, color=RGBColor(142, 191, 239), bold=True, align=PP_ALIGN.RIGHT)


def heading(slide, kicker, title_value, subtitle, page, footer_label):
    rect(slide, 0, 0, 13.333, 7.5, WHITE)
    rect(slide, 0, 0, 13.333, 0.14, BLUE)
    text(slide, kicker.upper(), 0.72, 0.42, 3.7, 0.2, size=10.5, color=BLUE, bold=True)
    text(slide, title_value, 0.7, 0.72, 11.9, 0.55, size=28, color=NAVY, bold=True, font="Aptos Display")
    text(slide, subtitle, 0.72, 1.38, 11.8, 0.34, size=13, color=MUTED)
    footer(slide, page, footer_label)


def card(slide, x, y, w, h, title_value, body, accent, *, fill=WHITE):
    rect(slide, x, y, w, h, fill, line=LINE, rounded=True)
    rect(slide, x, y, 0.09, h, accent, rounded=True)
    text(slide, title_value, x + 0.25, y + 0.22, w - 0.45, 0.28, size=16, color=accent, bold=True)
    text(slide, body, x + 0.25, y + 0.7, w - 0.5, h - 0.88, size=11.5, color=MUTED)


def step_card(slide, x, y, w, number, title_value, body, accent):
    rect(slide, x, y, w, 1.38, WHITE, line=LINE, rounded=True)
    rect(slide, x + 0.16, y + 0.17, 0.48, 0.48, accent, rounded=True)
    text(slide, number, x + 0.16, y + 0.29, 0.48, 0.16, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, title_value, x + 0.78, y + 0.18, w - 0.92, 0.24, size=14, color=NAVY, bold=True)
    text(slide, body, x + 0.78, y + 0.54, w - 0.92, 0.65, size=10, color=MUTED)


# 1 — Intro
slide = prs.slides.add_slide(BLANK)
rect(slide, 0, 0, 13.333, 7.5, NAVY)
rect(slide, 0, 0, 13.333, 0.16, ORANGE)
rect(slide, 8.72, 0.16, 4.613, 6.92, NAVY_2)
text(slide, "VendorLens AI", 0.84, 1.05, 7.25, 0.72, size=42, color=WHITE, bold=True, font="Aptos Display")
text(slide, "From supplier evidence to an\nauditable vendor decision", 0.87, 2.05, 7.4, 1.22, size=28, color=RGBColor(218, 232, 246), bold=True, font="Aptos Display")
text(slide, "AI-assisted • policy-driven • human-approved", 0.88, 3.65, 7.1, 0.28, size=16, color=ORANGE, bold=True)
text(slide, "As-built capstone presentation  |  October 2026", 0.88, 5.94, 6.8, 0.22, size=11, color=RGBColor(173, 198, 223))

for index, (label, note, accent) in enumerate([
    ("VALIDATE", "staged evidence", BLUE),
    ("GROUND", "RAG + policy", TEAL),
    ("DECIDE", "human review", VIOLET),
    ("CREATE", "Vendor Master", ORANGE),
]):
    y = 1.05 + index * 1.32
    rect(slide, 9.32, y, 3.0, 0.82, accent, rounded=True)
    text(slide, label, 9.54, y + 0.16, 1.15, 0.22, size=14, color=WHITE, bold=True)
    text(slide, note, 10.65, y + 0.18, 1.42, 0.2, size=10.5, color=WHITE, align=PP_ALIGN.RIGHT)
footer(slide, 1, "Introduction")


# 2 — Problem, cost, why AI
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Business case",
    "Why supplier onboarding needs a different operating model",
    "AI is useful where the work is repetitive and evidence-heavy—not where accountability is required.",
    2,
    "Problem • cost of status quo • why AI",
)

columns = [
    (0.72, "THE PROBLEM", "Fragmented evidence", [
        "Registration, tax, bank, insurance and policy evidence arrive through separate channels.",
        "Reviewers repeatedly search, transcribe and reconcile the same facts.",
        "Decisions are difficult to reproduce when source evidence is not linked.",
    ], RED, PALE_ORANGE),
    (4.52, "COST OF STATUS QUO", "Delay, rework and risk", [
        "Longer onboarding cycles delay supplier readiness and procurement activity.",
        "Wrong or incomplete files create avoidable back-and-forth.",
        "Manual judgement varies; unsupported approvals and privacy leakage become harder to detect.",
    ], ORANGE, RGBColor(255, 250, 241)),
    (8.32, "WHY AI", "Assist the reviewer", [
        "OCR and extraction turn unstructured evidence into reviewable fields.",
        "RAG finds relevant passages and returns source-linked answers.",
        "AI drafts; deterministic gates and the human reviewer make the decision.",
    ], BLUE, PALE_BLUE),
]
for x, kicker, title_value, items, accent, fill in columns:
    rect(slide, x, 2.05, 3.55, 4.42, fill, line=LINE, rounded=True)
    text(slide, kicker, x + 0.28, 2.34, 2.9, 0.2, size=10, color=accent, bold=True)
    text(slide, title_value, x + 0.28, 2.72, 2.95, 0.56, size=21, color=NAVY, bold=True, font="Aptos Display")
    bullets(slide, items, x + 0.28, 3.55, 2.95, 2.45, size=12, color=INK, gap=10)


# 3 — Application journey architecture
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Architecture 1/3 • journey",
    "A gated path from upload to Vendor Master",
    "Each boundary has a clear owner, failure mode and audit record.",
    3,
    "Application journey",
)

journey = [
    ("01", "Configure", "Category, supplier details and dynamic checklist", BLUE),
    ("02", "Validate", "Stage file • OCR • type, field and identity checks", TEAL),
    ("03", "Process", "Redact • extract • embed • evaluate policy", VIOLET),
    ("04", "Review", "Compare originals • correct • confirm or flag", ORANGE),
    ("05", "Create", "Approve • ERP record • 10-digit Vendor ID", GREEN),
]
for index, (number, title_value, body, accent) in enumerate(journey):
    x = 0.62 + index * 2.53
    step_card(slide, x, 2.2, 2.18, number, title_value, body, accent)
    if index < len(journey) - 1:
        text(slide, "→", x + 2.22, 2.59, 0.28, 0.24, size=20, color=MUTED, bold=True, align=PP_ALIGN.CENTER)

rect(slide, 0.78, 3.88, 11.78, 0.1, LINE)
text(slide, "Key control points", 0.8, 4.22, 2.1, 0.24, size=16, color=NAVY, bold=True)
controls = [
    (0.8, "Rejected uploads never become evidence", "Unreadable, wrong-type, incomplete, mismatched or poor OCR files are removed.", RED),
    (4.38, "Manual-attention items stay visible", "Uncertain OCR values are previewed read-only and excluded from safe bulk confirmation.", ORANGE),
    (7.96, "Approval requires explicit action", "All persisted compliance results must pass before ERP creation and case finalization.", GREEN),
]
for x, title_value, body, accent in controls:
    card(slide, x, 4.72, 3.25, 1.38, title_value, body, accent, fill=PALE)


# 4 — Full-stack architecture
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Architecture 2/3 • full stack",
    "A bounded full-stack system with explicit data ownership",
    "PostgreSQL is authoritative; vectors and files support evidence processing; ERP remains a separate boundary.",
    4,
    "Full-stack architecture",
)

# Role layer
text(slide, "EXPERIENCE", 0.55, 2.12, 1.1, 0.2, size=9.5, color=MUTED, bold=True)
for x, label, note, accent in [
    (1.75, "Supplier portal", "details • uploads • OCR preview", BLUE),
    (5.03, "Reviewer workbench", "evidence • checks • decision", VIOLET),
    (8.31, "Admin + Vendor Master", "operations • ERP records", ORANGE),
]:
    rect(slide, x, 1.98, 2.8, 0.82, accent, rounded=True)
    text(slide, label, x, 2.14, 2.8, 0.2, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, note, x + 0.1, 2.46, 2.6, 0.16, size=9, color=WHITE, align=PP_ALIGN.CENTER)

text(slide, "API + DOMAIN", 0.55, 3.37, 1.25, 0.2, size=9.5, color=MUTED, bold=True)
rect(slide, 1.75, 3.17, 9.36, 0.94, NAVY, rounded=True)
text(slide, "FastAPI", 2.05, 3.42, 1.15, 0.24, size=17, color=WHITE, bold=True)
text(slide, "auth • staged ingestion • AI orchestration • compliance • audit • ERP client", 3.25, 3.43, 7.45, 0.22, size=12.5, color=RGBColor(218, 232, 246), align=PP_ALIGN.CENTER)

text(slide, "DATA + SERVICES", 0.55, 4.75, 1.35, 0.2, size=9.5, color=MUTED, bold=True)
services = [
    (1.75, "PostgreSQL", "workflow + audit", BLUE),
    (4.12, "Upload volume", "original evidence", TEAL),
    (6.49, "ChromaDB", "redacted vectors", VIOLET),
    (8.86, "AI provider", "OpenRouter / Azure", ORANGE),
    (11.23, "Mock ERP MCP", "Vendor Master", GREEN),
]
for x, title_value, note, accent in services:
    rect(slide, x, 4.54, 1.9, 0.94, PALE, line=accent, rounded=True)
    text(slide, title_value, x, 4.73, 1.9, 0.2, size=13, color=accent, bold=True, align=PP_ALIGN.CENTER)
    text(slide, note, x + 0.06, 5.08, 1.78, 0.16, size=9, color=MUTED, align=PP_ALIGN.CENTER)

rect(slide, 1.75, 5.92, 11.38, 0.58, PALE_BLUE, line=RGBColor(188, 214, 239), rounded=True)
text(slide, "Cross-cutting: role authorization • PII redaction • integrity checks • Langfuse traces • Prometheus metrics", 1.98, 6.1, 10.92, 0.18, size=11.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


# 5 — RAG and decision architecture
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Architecture 3/3 • RAG",
    "Grounded answers without surrendering the decision",
    "Document RAG answers case questions; policy grounding assists review and correction messaging.",
    5,
    "RAG + human-in-the-loop controls",
)

rag_steps = [
    (0.62, "1", "Redact", "GSTIN, PAN, CIN, email, phone and bank values become typed placeholders.", RED),
    (3.05, "2", "Chunk", "Page-local 500-token chunks with 75-token overlap preserve citations.", BLUE),
    (5.48, "3", "Retrieve", "Embed the question; filter by supplier before top-K similarity ranking.", VIOLET),
    (7.91, "4", "Generate", "Answer only from retrieved evidence and return exact supporting chunk labels.", TEAL),
    (10.34, "5", "Verify", "Validate citations or return the guarded not-found response.", GREEN),
]
for x, number, title_value, body, accent in rag_steps:
    rect(slide, x, 2.1, 2.16, 2.1, WHITE, line=LINE, rounded=True)
    rect(slide, x + 0.17, 2.28, 0.43, 0.43, accent, rounded=True)
    text(slide, number, x + 0.17, 2.39, 0.43, 0.15, size=10.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, title_value, x + 0.72, 2.3, 1.22, 0.24, size=15, color=accent, bold=True)
    text(slide, body, x + 0.2, 2.94, 1.76, 0.93, size=10.2, color=MUTED)

rect(slide, 0.75, 4.72, 5.85, 1.35, PALE_TEAL, line=RGBColor(175, 224, 215), rounded=True)
text(slide, "Document RAG", 1.02, 4.98, 5.15, 0.25, size=16, color=TEAL, bold=True)
text(slide, "Reviewer questions → supplier-filtered evidence → cited answer / exact not-found", 1.02, 5.43, 5.15, 0.28, size=11.5, color=INK)

rect(slide, 6.78, 4.72, 5.8, 1.35, PALE_ORANGE, line=RGBColor(239, 203, 176), rounded=True)
text(slide, "Policy + review grounding", 7.05, 4.98, 5.05, 0.25, size=16, color=ORANGE, bold=True)
text(slide, "Policy checks and AI draft reasons remain editable; bulk confirmation excludes manual-attention items.", 7.05, 5.43, 5.05, 0.36, size=11.5, color=INK)

text(slide, "AI proposes evidence-backed output  →  deterministic gates expose gaps  →  reviewer confirms the outcome", 1.1, 6.35, 11.15, 0.24, size=13.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)


# 6 — Demo
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Live demo",
    "Show one complete supplier journey",
    "The demo should emphasize visible controls and recovery paths—not only the happy path.",
    6,
    "Live product walkthrough",
)

demo_steps = [
    (0.75, 2.08, "01", "Register + configure", "Create the supplier, select the category and show the generated checklist.", BLUE),
    (4.55, 2.08, "02", "Upload + validate", "Show a rejection, then a valid scan with OCR reliability and read-only values.", TEAL),
    (8.35, 2.08, "03", "Submit + process", "Snapshot requirements; watch PROCESSING move to reviewer attention.", VIOLET),
    (0.75, 4.18, "04", "Review evidence", "Filter matched/attention items, open originals and correct one extracted value.", ORANGE),
    (4.55, 4.18, "05", "Confirm or flag", "Use safe bulk confirmation; generate and edit a grounded correction reason.", RED),
    (8.35, 4.18, "06", "Approve + create", "Validate ERP payload; show portal reference, ERP record ID and Vendor ID.", GREEN),
]
for x, y, number, title_value, body, accent in demo_steps:
    rect(slide, x, y, 3.48, 1.62, WHITE, line=LINE, rounded=True)
    rect(slide, x + 0.2, y + 0.22, 0.52, 0.52, accent, rounded=True)
    text(slide, number, x + 0.2, y + 0.36, 0.52, 0.16, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, title_value, x + 0.88, y + 0.24, 2.3, 0.25, size=15, color=NAVY, bold=True)
    text(slide, body, x + 0.88, y + 0.67, 2.22, 0.68, size=10.5, color=MUTED)

text(slide, "Finish in the read-only Vendor Master: search the new Vendor ID, retrieve the ERP record and link back to onboarding evidence.", 1.0, 6.35, 11.3, 0.27, size=12.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


# 7 — Metrics placeholder
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Metrics • placeholder",
    "Measurement framework ready; results to be inserted",
    "Keep this slide value-free until the final monitored validation run is complete.",
    7,
    "Metrics placeholder",
)

metric_cards = [
    (0.72, "QUALITY", "TBD", "Extraction accuracy\nGrounding + citation rate\nNot-found + isolation", BLUE),
    (3.85, "EFFICIENCY", "TBD", "Median review time\nTouchless validation rate\nCorrection cycle count", TEAL),
    (6.98, "RELIABILITY", "TBD", "Upload/AI success rate\nP95 processing latency\nERP retry success", ORANGE),
    (10.11, "GOVERNANCE", "TBD", "Manual-attention rate\nReviewer overrides\nAudit completeness", VIOLET),
]
for x, label, value, body, accent in metric_cards:
    item = rect(slide, x, 2.12, 2.48, 3.03, PALE, line=accent, rounded=True)
    item.line.width = Pt(1.5)
    text(slide, label, x + 0.22, 2.4, 2.02, 0.2, size=10.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    text(slide, value, x + 0.22, 2.95, 2.02, 0.55, size=29, color=NAVY, bold=True, align=PP_ALIGN.CENTER, font="Aptos Display")
    text(slide, body, x + 0.28, 3.77, 1.9, 0.88, size=11, color=MUTED, align=PP_ALIGN.CENTER)

rect(slide, 1.04, 5.63, 11.2, 0.7, WHITE, line=LINE, rounded=True)
text(slide, "Data sources", 1.3, 5.85, 1.2, 0.2, size=11, color=NAVY, bold=True)
text(slide, "Langfuse traces + scores   •   Prometheus service metrics   •   Promptfoo/golden-set evaluation   •   reviewer audit events", 2.65, 5.83, 9.2, 0.22, size=11.5, color=MUTED, align=PP_ALIGN.CENTER)


# 8 — Limitations, trade-offs, next steps
slide = prs.slides.add_slide(BLANK)
heading(
    slide,
    "Close",
    "Limitations, trade-offs and the path to production",
    "The prototype proves the workflow; production readiness requires stronger identity, storage, operations and governance.",
    8,
    "Limitations • trade-offs • next steps",
)

sections = [
    (0.72, "LIMITATIONS", [
        "Synthetic India-only policy and mock ERP",
        "Local file/vector storage and in-process background work",
        "Demo authentication and no malware scanning",
        "OCR and model outputs remain probabilistic",
    ], RED, PALE_ORANGE),
    (4.52, "TRADE-OFFS", [
        "Pre-upload AI gates improve quality but add latency and provider dependency",
        "Human confirmation preserves accountability but not full automation",
        "Local OCR protects evidence flow but consumes application resources",
        "ChromaDB simplifies the demo but is not the final scale design",
    ], ORANGE, RGBColor(255, 250, 241)),
    (8.32, "NEXT STEPS", [
        "Enterprise SSO/RBAC, TLS and secrets management",
        "Encrypted object storage, malware scanning and retention controls",
        "Durable worker queue, retries and reconciliation",
        "Production ERP adapter plus final evaluation and monitoring gates",
    ], GREEN, PALE_TEAL),
]
for x, label, items, accent, fill in sections:
    rect(slide, x, 2.04, 3.55, 4.12, fill, line=LINE, rounded=True)
    text(slide, label, x + 0.28, 2.35, 2.95, 0.22, size=11, color=accent, bold=True)
    bullets(slide, items, x + 0.28, 2.93, 2.92, 2.75, size=11.5, color=INK, gap=10)

rect(slide, 1.45, 6.37, 10.45, 0.42, NAVY, rounded=True)
text(slide, "North star: reduce reviewer effort while keeping evidence, policy and accountability visible.", 1.65, 6.49, 10.05, 0.16, size=11.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


prs.core_properties.title = "VendorLens AI - Presentation Deck"
prs.core_properties.subject = "Supplier onboarding, full-stack architecture, RAG, demo, metrics, and production path"
prs.core_properties.author = "VendorLens AI"
prs.core_properties.comments = "As-built baseline: merge_branch c2a23e4"
prs.save(OUT)
print(OUT)
